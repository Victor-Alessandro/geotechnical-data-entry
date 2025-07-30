"""
Originally I did a similar substitution to each docx file so they could be changed with media replacement. The word placeholder images were thus effectively a GUI and whatever changes in format were imposed on them would also be reflected once the template was populated. However, I lost the original and made a mistake on the model (it was a single one as well). There are multiple steps to reproduce previous results: Preparing the template and placeholders, applying the inlining and then performing media replacement during document generation. I do not have the time to identify the bug within these steps and the gain would be minimal.

By the way, the way pasted images are represented on the online and offline versions of word are different. So only the offline one should be used. the online version is a second place citizen.
"""


import os
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm

# --- Configuration ---
BASE_TEMPLATE_PATH = "aciers.docx"
PREPARED_TEMPLATE_PATH = "aciers_prepared.docx"
PLACEHOLDER_IMAGE_FOLDER = "placeholders"

# This dictionary maps the Jinja2 tag in your document to the placeholder image filename.
# The keys MUST match the tags in your Word template exactly.
IMAGE_MAPPING = {
    "Placeholder_1": "Placeholder_1.png",
    "Placeholder_3": "Placeholder_3.png",
    "Placeholder_4": "Placeholder_4.png",
    "img_beton": "img_beton.png",
}

# --- Script Execution ---
def prepare_template_with_images():
    """
    Renders a base template with placeholder images to create a new, 
    prepared template for later use with replace_media.
    """
    if not os.path.exists(BASE_TEMPLATE_PATH):
        print(f"Error: Base template '{BASE_TEMPLATE_PATH}' not found.")
        return

    doc = DocxTemplate(BASE_TEMPLATE_PATH)
    context = {}

    print("Preparing context with placeholder images...")
    for tag, filename in IMAGE_MAPPING.items():
        image_path = os.path.join(PLACEHOLDER_IMAGE_FOLDER, filename)
        
        if os.path.exists(image_path):
            # Create an InlineImage object for the placeholder.
            # You can set a default size for your placeholders here.
            image = InlineImage(doc, image_path, width=Mm(80))
            context[tag] = image
            print(f"  - Mapping tag '{{{{ {tag} }}}}' to image '{image_path}'")
        else:
            print(f"  - WARNING: Image not found for tag '{{{{ {tag} }}}}' at '{image_path}'")
            context[tag] = f"[Image: {filename} not found]"

    # Render the document to insert the placeholder images
    doc.render(context)

    # Save the new document with the images embedded
    doc.save(PREPARED_TEMPLATE_PATH)
    
    print(f"\n✅ Successfully created prepared template: '{PREPARED_TEMPLATE_PATH}'")
    print("This file now contains embedded images and is ready for use with 'replace_media'.")

if __name__ == "__main__":
    # Create dummy folder and files for demonstration if they don't exist
    if not os.path.exists(PLACEHOLDER_IMAGE_FOLDER):
        os.makedirs(PLACEHOLDER_IMAGE_FOLDER)
        print(f"Created dummy folder: '{PLACEHOLDER_IMAGE_FOLDER}'")
        for fname in IMAGE_MAPPING.values():
            with open(os.path.join(PLACEHOLDER_IMAGE_FOLDER, fname), "w") as f:
                f.write("dummy")

    prepare_template_with_images()
